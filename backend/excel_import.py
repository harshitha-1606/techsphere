"""Excel parsing + validation. Columns are detected by header name, never by position."""
import datetime as dt
import io
import re

from . import config

EMAIL_RE = re.compile(r"^[A-Za-z0-9._%+\-']+@[A-Za-z0-9\-]+(\.[A-Za-z0-9\-]+)*\.[A-Za-z]{2,}$")

# canonical field -> accepted normalised header names (lower-case, letters/digits only)
SYNONYMS = {
    "department": {"department", "dept", "branch", "departmentname"},
    "team": {"teambatch", "team", "batch", "teamname", "batchname"},
    "name": {"studentname", "name", "participantname", "nameofthestudent", "student"},
    "roll_no": {"rollno", "rollnumber", "rollnos", "regno", "registrationnumber", "regdno"},
    "email": {"emailid", "email", "emailaddress", "mailid", "mail", "gmail"},
    "mobile": {"mobilenumber", "mobileno", "mobile", "phone", "phonenumber", "contactnumber"},
    "topic": {"themetopic", "theme", "topic", "posterthemetopic", "posterthemetitle", "title"},
}
REQUIRED = ("name", "department", "email")
LABELS = {"name": "Student Name", "department": "Department", "email": "Email ID",
          "team": "Team / Batch", "roll_no": "Roll No.", "mobile": "Mobile Number", "topic": "Theme / Topic"}


class ImportFileError(Exception):
    """User-facing problem with the uploaded file."""


def _norm_header(v):
    return re.sub(r"[^a-z0-9]", "", str(v or "").lower())


def _clean(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.datetime, dt.date)):
        v = v.isoformat()
    return re.sub(r"\s+", " ", str(v).replace("\u00a0", " ")).strip()


def _grids_xlsx(content):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise ImportFileError("The file could not be read as an .xlsx workbook. It may be corrupted.") from exc
    for ws in wb.worksheets:
        grid = [[c for c in row] for row in ws.iter_rows(values_only=True)]
        for rng in ws.merged_cells.ranges:  # fill merged cells (e.g. a merged Department cell)
            val = ws.cell(rng.min_row, rng.min_col).value
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    if r - 1 < len(grid) and c - 1 < len(grid[r - 1]):
                        grid[r - 1][c - 1] = val
        yield ws.title, grid


def _grids_xls(content):
    try:
        import xlrd
    except ImportError as exc:
        raise ImportFileError("Reading .xls needs the 'xlrd' package. Install it or save the file as .xlsx.") from exc
    try:
        wb = xlrd.open_workbook(file_contents=content, formatting_info=True)
    except Exception:
        try:
            wb = xlrd.open_workbook(file_contents=content)
        except Exception as exc:
            raise ImportFileError("The file could not be read as an .xls workbook.") from exc
    for sh in wb.sheets():
        grid = [[sh.cell_value(r, c) if sh.cell_type(r, c) not in (0, 6) else None
                 for c in range(sh.ncols)] for r in range(sh.nrows)]
        for (r0, r1, c0, c1) in getattr(sh, "merged_cells", []):
            val = grid[r0][c0]
            for r in range(r0, r1):
                for c in range(c0, c1):
                    grid[r][c] = val
        yield sh.name, grid


def _find_header(grid):
    """Return (row_index, {field: col_index}) for the first row that contains all required headers."""
    best = None
    for i, row in enumerate(grid[:25]):
        cols = {}
        for j, cell in enumerate(row):
            n = _norm_header(cell)
            if not n:
                continue
            for field, names in SYNONYMS.items():
                if n in names and field not in cols:
                    cols[field] = j
        if all(f in cols for f in REQUIRED):
            return i, cols
        if best is None or len(cols) > len(best[1]):
            best = (i, cols)
    return None, (best[1] if best else {})


def parse_workbook(content, filename):
    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        grids = _grids_xlsx(content)
    elif name.endswith(".xls"):
        grids = _grids_xls(content)
    else:
        raise ImportFileError("Only .xlsx and .xls files are accepted.")

    last_cols = {}
    for sheet, grid in grids:
        hdr, cols = _find_header(grid)
        last_cols = cols or last_cols
        if hdr is None:
            continue
        return _validate(grid, hdr, cols, sheet)
    missing = [LABELS[f] for f in REQUIRED if f not in last_cols]
    raise ImportFileError("Required column(s) not found: " + ", ".join(missing or ["header row"]) +
                          ". Expected a header row with Student Name, Department and Email ID.")


def _validate(grid, hdr, cols, sheet):
    rows, blank = [], 0
    seen_email, seen_key = {}, {}
    data_rows = grid[hdr + 1:]
    if len(data_rows) > config.MAX_ROWS:
        raise ImportFileError(f"Too many rows ({len(data_rows)}). The limit is {config.MAX_ROWS}.")
    for offset, raw in enumerate(data_rows):
        excel_row = hdr + 2 + offset
        rec = {f: _clean(raw[j]) if j < len(raw) else "" for f, j in cols.items()}
        for f in SYNONYMS:
            rec.setdefault(f, "")
        if not any(rec.values()):
            blank += 1
            continue
        errors, warnings = [], []
        rec["email"] = rec["email"].replace(" ", "").lower()
        for f in REQUIRED:
            if not rec[f]:
                errors.append(f"Missing {LABELS[f]}")
        if rec["email"] and not EMAIL_RE.match(rec["email"]):
            errors.append("Invalid email format")
        if len(rec["name"]) > config.MAX_NAME_LEN:
            errors.append(f"Student Name longer than {config.MAX_NAME_LEN} characters")
        if len(rec["department"]) > config.MAX_DEPT_LEN:
            errors.append(f"Department longer than {config.MAX_DEPT_LEN} characters")
        if rec["email"]:
            if rec["email"] in seen_email:
                errors.append(f"Duplicate email (first used in row {seen_email[rec['email']]})")
            else:
                seen_email[rec["email"]] = excel_row
        key = (rec["name"].lower(), rec["department"].lower(), rec["roll_no"].lower())
        if rec["name"] and rec["department"]:
            if key in seen_key and rec["email"] != seen_key[key][1]:
                warnings.append(f"Same name/department/roll no. as row {seen_key[key][0]} with a different email")
            elif key in seen_key and not errors:
                errors.append(f"Duplicate certificate record (same as row {seen_key[key][0]})")
            seen_key.setdefault(key, (excel_row, rec["email"]))
        rows.append({"row": excel_row, "data": rec, "errors": errors, "warnings": warnings,
                     "valid": not errors})
    valid = sum(1 for r in rows if r["valid"])
    return {
        "sheet": sheet,
        "columns_detected": {LABELS[f]: bool(f in cols) for f in LABELS},
        "total": len(rows), "valid": valid, "invalid": len(rows) - valid,
        "blank_rows_skipped": blank, "rows": rows,
    }
